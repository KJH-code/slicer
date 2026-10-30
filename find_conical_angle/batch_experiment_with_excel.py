"""Batch runner with per-model Excel reports and one sheet per slice band.

This module reuses the stable batch-analysis helpers from ``batch_experiment``
and adds Excel generation.  Each model workbook contains the configured
measurement sheets plus ``밴드1``, ``밴드2``, ... up to the maximum band count
observed for that model.

Band-sheet layout:
- rows: q0
- columns: kappa_max
- cell value: mean signed band angle across successful repeats
  (inward is negative, outward is positive)
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import sys
import time
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from batch_experiment import (
    BATCH_DEFAULTS as CORE_BATCH_DEFAULTS,
    SUMMARY_FIELDS as CORE_SUMMARY_FIELDS,
    BatchConfigError,
    _default_run_name,
    _expand_grid_values,
    _read_json_object,
    _resolve_path,
    _safe_filename,
    _utc_now,
    _write_jsonl_row,
    _write_manifest,
    discover_stl_files,
    error_to_summary as core_error_to_summary,
    expand_experiment_configs,
    result_to_summary as core_result_to_summary,
)
from conical_slicing import analyze_stl
from conical_slicing.analysis import write_json
from conical_slicing.config import DEFAULT_CONFIG, ConfigError
from conical_slicing.mesh import MeshAnalysisError


EXCEL_MEASUREMENT_SPECS: dict[str, dict[str, str]] = {
    "selected_angle": {
        "sheet": "선택원뿔각",
        "title": "선택 원뿔각 (deg)",
        "number_format": "0.000",
    },
    "runtime_seconds": {
        "sheet": "분석시간",
        "title": "분석 함수 실행 시간 (s)",
        "number_format": "0.000000",
    },
    "wall_runtime_seconds": {
        "sheet": "전체경과시간",
        "title": "전체 경과 시간 (s)",
        "number_format": "0.000000",
    },
    "selected_loss": {
        "sheet": "선택손실",
        "title": "선택 결과 손실값",
        "number_format": "0.000000",
    },
    "selected_direction": {
        "sheet": "선택방향",
        "title": "선택 방향",
        "number_format": "@",
    },
    "selected_violating_area": {
        "sheet": "위반면적",
        "title": "선택 각도에서의 위반 면적",
        "number_format": "0.000000",
    },
    "selected_passed_area": {
        "sheet": "통과면적",
        "title": "선택 각도에서의 통과 면적",
        "number_format": "0.000000",
    },
    "selected_infeasible_face_count": {
        "sheet": "위반면수",
        "title": "선택 각도에서의 비허용 면 개수",
        "number_format": "0",
    },
    "warning_count": {
        "sheet": "경고수",
        "title": "메시 경고 개수",
        "number_format": "0",
    },
}

DEFAULT_EXCEL_MEASUREMENTS = list(EXCEL_MEASUREMENT_SPECS)

BATCH_DEFAULTS: dict[str, Any] = {
    **CORE_BATCH_DEFAULTS,
    "save_model_excel": True,
    "excel_include_raw_data": True,
    "excel_measurements": DEFAULT_EXCEL_MEASUREMENTS,
}

BAND_SUMMARY_FIELDS = ["slice_band_count", "slice_bands_json"]
SUMMARY_FIELDS = list(CORE_SUMMARY_FIELDS)
_error_type_index = (
    SUMMARY_FIELDS.index("error_type")
    if "error_type" in SUMMARY_FIELDS
    else len(SUMMARY_FIELDS)
)
for _offset, _field in enumerate(BAND_SUMMARY_FIELDS):
    if _field not in SUMMARY_FIELDS:
        SUMMARY_FIELDS.insert(_error_type_index + _offset, _field)


def load_batch_config(path: str | Path) -> tuple[dict[str, Any], Path]:
    """Load batch and Excel settings.

    Relative paths are interpreted from the directory containing the batch
    configuration file.  The original JSON format remains valid because all
    Excel settings have defaults.
    """

    config_path = Path(path).expanduser().resolve()
    raw = _read_json_object(config_path)
    unknown = sorted(set(raw) - set(BATCH_DEFAULTS))
    if unknown:
        raise BatchConfigError(f"알 수 없는 배치 설정 항목: {', '.join(unknown)}")

    config = dict(BATCH_DEFAULTS)
    config.update(raw)

    for key in (
        "recursive",
        "continue_on_error",
        "save_detailed_json",
        "save_model_excel",
        "excel_include_raw_data",
    ):
        if not isinstance(config[key], bool):
            raise BatchConfigError(f"{key}는 true 또는 false여야 합니다.")

    repeats = config["repeats"]
    if isinstance(repeats, bool) or not isinstance(repeats, int) or repeats <= 0:
        raise BatchConfigError("repeats는 1 이상의 정수여야 합니다.")

    run_name = config["run_name"]
    if run_name is not None and (
        not isinstance(run_name, str) or not run_name.strip()
    ):
        raise BatchConfigError(
            "run_name은 null 또는 비어 있지 않은 문자열이어야 합니다."
        )

    for key in ("models_dir", "output_dir", "base_config_file"):
        if not isinstance(config[key], str) or not config[key].strip():
            raise BatchConfigError(f"{key}는 비어 있지 않은 문자열이어야 합니다.")

    if not isinstance(config["base_overrides"], dict):
        raise BatchConfigError("base_overrides는 JSON 객체여야 합니다.")
    if not isinstance(config["parameter_grid"], dict):
        raise BatchConfigError("parameter_grid는 JSON 객체여야 합니다.")
    if not isinstance(config["parameter_sets"], list) or any(
        not isinstance(item, dict) for item in config["parameter_sets"]
    ):
        raise BatchConfigError("parameter_sets는 JSON 객체들의 배열이어야 합니다.")

    excel_measurements = config["excel_measurements"]
    if (
        not isinstance(excel_measurements, list)
        or not excel_measurements
        or any(
            not isinstance(item, str) or not item.strip()
            for item in excel_measurements
        )
    ):
        raise BatchConfigError(
            "excel_measurements는 비어 있지 않은 문자열 배열이어야 합니다."
        )

    unknown_measurements = sorted(
        set(excel_measurements) - set(SUMMARY_FIELDS)
    )
    if unknown_measurements:
        raise BatchConfigError(
            "excel_measurements의 알 수 없는 요약 필드: "
            + ", ".join(unknown_measurements)
        )

    for key, specification in config["parameter_grid"].items():
        if key not in DEFAULT_CONFIG:
            raise BatchConfigError(
                f"parameter_grid의 알 수 없는 설정 항목: {key}"
            )
        _expand_grid_values(key, specification)

    for index, parameter_set in enumerate(config["parameter_sets"], start=1):
        unknown_keys = sorted(set(parameter_set) - set(DEFAULT_CONFIG))
        if unknown_keys:
            raise BatchConfigError(
                f"parameter_sets[{index}]의 알 수 없는 설정 항목: "
                f"{', '.join(unknown_keys)}"
            )

    base_dir = config_path.parent
    config["models_dir"] = _resolve_path(config["models_dir"], base_dir)
    config["output_dir"] = _resolve_path(config["output_dir"], base_dir)
    config["base_config_file"] = _resolve_path(
        config["base_config_file"], base_dir
    )
    return config, config_path


def _normalized_slice_bands(result: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Return JSON-safe slice-band records in workbook order."""

    slice_plan = result.get("slice_plan")
    if not isinstance(slice_plan, Mapping):
        return []

    raw_bands = slice_plan.get("bands")
    if not isinstance(raw_bands, list):
        return []

    bands: list[dict[str, Any]] = []
    for fallback_index, raw in enumerate(raw_bands):
        if not isinstance(raw, Mapping):
            continue

        direction = str(raw.get("direction", ""))
        angle = _float_or_none(raw.get("angle"))
        signed_angle = _float_or_none(raw.get("signed_angle"))
        if signed_angle is None and angle is not None:
            if direction == "inward":
                signed_angle = -abs(angle)
            elif direction == "outward":
                signed_angle = abs(angle)

        component_ids = raw.get("component_ids", [])
        if not isinstance(component_ids, list):
            component_ids = []

        band_id_value = raw.get("band_id", fallback_index)
        try:
            band_id = int(band_id_value)
        except (TypeError, ValueError):
            band_id = fallback_index

        bands.append(
            {
                "band_id": band_id,
                "direction": direction,
                "angle": angle,
                "signed_angle": signed_angle,
                "h_min": _float_or_none(raw.get("h_min")),
                "h_max": _float_or_none(raw.get("h_max")),
                "model_loss": _float_or_none(raw.get("model_loss")),
                "component_ids": [
                    int(value)
                    for value in component_ids
                    if isinstance(value, int)
                    and not isinstance(value, bool)
                ],
            }
        )

    # ``slice_plan["bands"]`` is normally already ordered by transformed height.
    # Explicit sorting makes the workbook stable if a producer returns it unsorted.
    bands.sort(
        key=lambda item: (
            float("inf") if item["h_min"] is None else item["h_min"],
            float("inf") if item["h_max"] is None else item["h_max"],
            item["band_id"],
        )
    )
    return bands


def result_to_summary(
    result: Mapping[str, Any],
    *,
    run_name: str,
    experiment_index: int,
    configuration_index: int,
    repeat: int,
    model_path: Path,
    detail_json: Path | None,
    wall_runtime_seconds: float,
    config: Mapping[str, Any],
) -> dict[str, Any]:
    row = core_result_to_summary(
        result,
        run_name=run_name,
        experiment_index=experiment_index,
        configuration_index=configuration_index,
        repeat=repeat,
        model_path=model_path,
        detail_json=detail_json,
        wall_runtime_seconds=wall_runtime_seconds,
        config=config,
    )
    bands = _normalized_slice_bands(result)
    row["slice_band_count"] = len(bands)
    row["slice_bands_json"] = json.dumps(
        bands,
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    )
    return row


def error_to_summary(
    exc: Exception,
    *,
    run_name: str,
    experiment_index: int,
    configuration_index: int,
    repeat: int,
    model_path: Path,
    wall_runtime_seconds: float,
    config: Mapping[str, Any],
) -> dict[str, Any]:
    row = core_error_to_summary(
        exc,
        run_name=run_name,
        experiment_index=experiment_index,
        configuration_index=configuration_index,
        repeat=repeat,
        model_path=model_path,
        wall_runtime_seconds=wall_runtime_seconds,
        config=config,
    )
    row["slice_band_count"] = ""
    row["slice_bands_json"] = ""
    return row


def _excel_sheet_name(value: str, used: set[str]) -> str:
    """Return an Excel-safe, unique sheet name."""

    cleaned = re.sub(r"[\\/*?:\[\]]+", "_", value).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used:
        suffix_text = f"_{suffix}"
        candidate = f"{cleaned[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _excel_report_filename(model_path: str, used: set[str]) -> str:
    path = Path(model_path)
    base = _safe_filename(path.stem, max_length=80)
    candidate = f"{base}.xlsx"
    suffix = 2
    while candidate.casefold() in used:
        candidate = f"{base}_{suffix}.xlsx"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _float_or_none(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    return number


def _parse_bands_json(row: Mapping[str, Any]) -> list[dict[str, Any]]:
    raw = row.get("slice_bands_json", "")
    if raw in (None, ""):
        return []
    try:
        value = json.loads(str(raw))
    except (json.JSONDecodeError, TypeError, ValueError):
        return []
    if not isinstance(value, list):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]


def _aggregate_excel_value(
    rows: Sequence[Mapping[str, Any]], field: str
) -> tuple[Any, bool, str]:
    """Aggregate repeated runs for one q0/kappa cell."""

    ok_rows = [
        row for row in rows if str(row.get("status", "")) == "ok"
    ]
    error_rows = [
        row for row in rows if str(row.get("status", "")) != "ok"
    ]
    source_rows = rows if field == "status" else ok_rows
    raw_values = [row.get(field, "") for row in source_rows]
    values = [value for value in raw_values if value not in (None, "")]

    comment_lines = [
        f"전체 반복: {len(rows)}회",
        f"성공: {len(ok_rows)}회",
        f"오류: {len(error_rows)}회",
    ]

    if not values:
        if error_rows:
            error_types = Counter(
                str(row.get("error_type", "오류")) for row in error_rows
            )
            summary = ", ".join(
                f"{name} {count}회"
                for name, count in error_types.most_common()
            )
            comment_lines.append(f"오류 종류: {summary}")
            return "ERROR", False, "\n".join(comment_lines)
        return None, False, "\n".join(comment_lines)

    numeric_values = [_float_or_none(value) for value in values]
    if all(value is not None for value in numeric_values):
        numbers = [
            float(value) for value in numeric_values if value is not None
        ]
        mean = math.fsum(numbers) / len(numbers)
        comment_lines.extend(
            [
                f"평균: {mean:.12g}",
                f"최솟값: {min(numbers):.12g}",
                f"최댓값: {max(numbers):.12g}",
            ]
        )
        return mean, True, "\n".join(comment_lines)

    text_values = [str(value) for value in values]
    counts = Counter(text_values)
    most_common_count = max(counts.values())
    most_common = sorted(
        value
        for value, count in counts.items()
        if count == most_common_count
    )[0]
    comment_lines.append(
        "값별 횟수: "
        + ", ".join(
            f"{value}={count}" for value, count in counts.most_common()
        )
    )
    return most_common, False, "\n".join(comment_lines)


def _aggregate_band_angle(
    rows: Sequence[Mapping[str, Any]],
    band_index: int,
) -> tuple[Any, bool, str]:
    """Aggregate one ordered band's signed angle for a q0/kappa cell."""

    ok_rows = [
        row for row in rows if str(row.get("status", "")) == "ok"
    ]
    error_rows = [
        row for row in rows if str(row.get("status", "")) != "ok"
    ]

    selected_bands: list[dict[str, Any]] = []
    missing_band_count = 0
    for row in ok_rows:
        bands = _parse_bands_json(row)
        if band_index < len(bands):
            selected_bands.append(bands[band_index])
        else:
            missing_band_count += 1

    comment_lines = [
        f"전체 반복: {len(rows)}회",
        f"성공: {len(ok_rows)}회",
        f"오류: {len(error_rows)}회",
        f"해당 밴드 존재: {len(selected_bands)}회",
        f"해당 밴드 없음: {missing_band_count}회",
    ]

    if not selected_bands:
        if error_rows and not ok_rows:
            error_types = Counter(
                str(row.get("error_type", "오류")) for row in error_rows
            )
            comment_lines.append(
                "오류 종류: "
                + ", ".join(
                    f"{name} {count}회"
                    for name, count in error_types.most_common()
                )
            )
            return "ERROR", False, "\n".join(comment_lines)
        return None, False, "\n".join(comment_lines)

    signed_angles = [
        value
        for band in selected_bands
        if (value := _float_or_none(band.get("signed_angle"))) is not None
    ]
    if not signed_angles:
        return None, False, "\n".join(comment_lines)

    mean = math.fsum(signed_angles) / len(signed_angles)
    comment_lines.extend(
        [
            f"부호 포함 각도 평균: {mean:.12g}",
            f"최솟값: {min(signed_angles):.12g}",
            f"최댓값: {max(signed_angles):.12g}",
        ]
    )

    directions = Counter(
        str(band.get("direction", ""))
        for band in selected_bands
        if str(band.get("direction", ""))
    )
    if directions:
        comment_lines.append(
            "방향: "
            + ", ".join(
                f"{name}={count}"
                for name, count in directions.most_common()
            )
        )

    h_min_values = [
        value
        for band in selected_bands
        if (value := _float_or_none(band.get("h_min"))) is not None
    ]
    h_max_values = [
        value
        for band in selected_bands
        if (value := _float_or_none(band.get("h_max"))) is not None
    ]
    if h_min_values:
        comment_lines.append(
            f"h_min 범위: {min(h_min_values):.12g} ~ "
            f"{max(h_min_values):.12g}"
        )
    if h_max_values:
        comment_lines.append(
            f"h_max 범위: {min(h_max_values):.12g} ~ "
            f"{max(h_max_values):.12g}"
        )

    component_sets = Counter(
        json.dumps(
            band.get("component_ids", []),
            ensure_ascii=False,
            separators=(",", ":"),
        )
        for band in selected_bands
    )
    if component_sets:
        comment_lines.append(
            "component_ids: "
            + ", ".join(
                f"{value}={count}"
                for value, count in component_sets.most_common()
            )
        )

    return mean, True, "\n".join(comment_lines)


def _prepare_matrix_sheet(
    worksheet: Any,
    *,
    title: str,
    description: str,
    q0_values: Sequence[float],
    kappa_values: Sequence[float],
    header_fill: Any,
    subheader_fill: Any,
    title_fill: Any,
    header_font: Any,
    title_font: Any,
    table_border: Any,
    center: Any,
) -> int:
    """Create common q0-by-kappa worksheet headers."""

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "B5"

    last_column = 1 + len(kappa_values)
    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.fill = title_fill
    title_cell.font = title_font
    title_cell.alignment = type(center)(
        horizontal="left",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 24

    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=last_column,
    )
    description_cell = worksheet.cell(
        row=2,
        column=1,
        value=description,
    )
    description_cell.alignment = type(center)(
        wrap_text=True,
        vertical="center",
    )
    worksheet.row_dimensions[2].height = 34

    corner = worksheet.cell(row=4, column=1, value="q0 \\ kappa_max")
    corner.fill = header_fill
    corner.font = header_font
    corner.alignment = center
    corner.border = table_border

    for column_index, kappa in enumerate(kappa_values, start=2):
        cell = worksheet.cell(row=4, column=column_index, value=kappa)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = table_border
        cell.number_format = "0.00"

    for row_index, q0 in enumerate(q0_values, start=5):
        q0_cell = worksheet.cell(row=row_index, column=1, value=q0)
        q0_cell.fill = subheader_fill
        q0_cell.font = type(header_font)(bold=True)
        q0_cell.alignment = center
        q0_cell.border = table_border
        q0_cell.number_format = "0.00"

    worksheet.column_dimensions["A"].width = 16
    return last_column


def write_model_excel_reports(
    summary_csv: Path,
    output_dir: Path,
    *,
    measurements: Sequence[str],
    include_raw_data: bool,
) -> list[Path]:
    """Create one workbook per model, including one sheet per slice band."""

    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import (
            Alignment,
            Border,
            Font,
            PatternFill,
            Side,
        )
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise BatchConfigError(
            "모델별 Excel 생성을 위해 openpyxl이 필요합니다. "
            "'py -m pip install openpyxl'을 실행한 뒤 다시 시도하세요."
        ) from exc

    with summary_csv.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as stream:
        rows = list(csv.DictReader(stream))

    if not rows:
        return []

    grouped_by_model: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        model_key = row.get("model_path") or row.get("model") or "model"
        grouped_by_model[str(model_key)].append(row)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_paths: list[Path] = []
    used_filenames: set[str] = set()

    header_fill = PatternFill("solid", fgColor="1F4E78")
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")
    title_fill = PatternFill("solid", fgColor="DDEBF7")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=14, bold=True, color="17365D")
    thin_gray = Side(style="thin", color="B7C9D6")
    table_border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )
    center = Alignment(horizontal="center", vertical="center")

    for model_path, model_rows in sorted(
        grouped_by_model.items(),
        key=lambda item: item[0].casefold(),
    ):
        model_rows.sort(
            key=lambda row: int(row.get("experiment_index") or 0)
        )

        q0_values = sorted(
            {
                value
                for row in model_rows
                if (
                    value := _float_or_none(row.get("config_q0"))
                )
                is not None
            }
        )
        kappa_values = sorted(
            {
                value
                for row in model_rows
                if (
                    value := _float_or_none(
                        row.get("config_kappa_max")
                    )
                )
                is not None
            }
        )
        if not q0_values or not kappa_values:
            print(
                f"경고: {Path(model_path).name}의 q0 또는 "
                "kappa_max 값이 없어 Excel 보고서를 건너뜁니다.",
                file=sys.stderr,
            )
            continue

        rows_by_cell: dict[
            tuple[float, float],
            list[dict[str, str]],
        ] = defaultdict(list)
        for row in model_rows:
            q0 = _float_or_none(row.get("config_q0"))
            kappa = _float_or_none(row.get("config_kappa_max"))
            if q0 is not None and kappa is not None:
                rows_by_cell[(q0, kappa)].append(row)

        maximum_band_count = max(
            (
                len(_parse_bands_json(row))
                for row in model_rows
                if str(row.get("status", "")) == "ok"
            ),
            default=0,
        )

        workbook = Workbook()
        workbook.remove(workbook.active)
        used_sheet_names: set[str] = set()

        for field in measurements:
            spec = EXCEL_MEASUREMENT_SPECS.get(
                field,
                {
                    "sheet": field,
                    "title": field,
                    "number_format": "0.000000",
                },
            )
            sheet_name = _excel_sheet_name(
                spec["sheet"],
                used_sheet_names,
            )
            worksheet = workbook.create_sheet(sheet_name)

            last_column = _prepare_matrix_sheet(
                worksheet,
                title=spec["title"],
                description=(
                    "행: q0 / 열: kappa_max / 숫자: 반복 평균 / "
                    "문자열: 반복 최빈값. 셀 메모에 반복 통계가 "
                    "저장됩니다."
                ),
                q0_values=q0_values,
                kappa_values=kappa_values,
                header_fill=header_fill,
                subheader_fill=subheader_fill,
                title_fill=title_fill,
                header_font=header_font,
                title_font=title_font,
                table_border=table_border,
                center=center,
            )

            numeric_sheet = True
            for row_index, q0 in enumerate(q0_values, start=5):
                for column_index, kappa in enumerate(
                    kappa_values,
                    start=2,
                ):
                    value, is_numeric, comment = (
                        _aggregate_excel_value(
                            rows_by_cell.get((q0, kappa), []),
                            field,
                        )
                    )
                    cell = worksheet.cell(
                        row=row_index,
                        column=column_index,
                        value=value,
                    )
                    cell.alignment = center
                    cell.border = table_border
                    cell.number_format = (
                        spec["number_format"] if is_numeric else "@"
                    )
                    cell.comment = Comment(
                        comment,
                        "batch_experiment",
                    )
                    if value is not None and not is_numeric:
                        numeric_sheet = False

            for column_index in range(2, last_column + 1):
                worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = 12

            if numeric_sheet:
                data_range = (
                    f"B5:{get_column_letter(last_column)}"
                    f"{4 + len(q0_values)}"
                )
                worksheet.conditional_formatting.add(
                    data_range,
                    ColorScaleRule(
                        start_type="min",
                        start_color="FEE2E2",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FEF3C7",
                        end_type="max",
                        end_color="DCFCE7",
                    ),
                )

        # One extra worksheet for each ordered band observed in this model.
        for band_index in range(maximum_band_count):
            worksheet = workbook.create_sheet(
                _excel_sheet_name(
                    f"밴드{band_index + 1}",
                    used_sheet_names,
                )
            )
            last_column = _prepare_matrix_sheet(
                worksheet,
                title=(
                    f"밴드 {band_index + 1} 원뿔각 "
                    "(deg, inward 음수 / outward 양수)"
                ),
                description=(
                    "행: q0 / 열: kappa_max / 셀: 반복 실행에서의 "
                    "부호 포함 원뿔각 평균. 밴드 순서는 변환 높이 "
                    "h_min 기준이며, 해당 조합에 이 밴드가 없으면 "
                    "빈 셀입니다."
                ),
                q0_values=q0_values,
                kappa_values=kappa_values,
                header_fill=header_fill,
                subheader_fill=subheader_fill,
                title_fill=title_fill,
                header_font=header_font,
                title_font=title_font,
                table_border=table_border,
                center=center,
            )

            for row_index, q0 in enumerate(q0_values, start=5):
                for column_index, kappa in enumerate(
                    kappa_values,
                    start=2,
                ):
                    value, is_numeric, comment = (
                        _aggregate_band_angle(
                            rows_by_cell.get((q0, kappa), []),
                            band_index,
                        )
                    )
                    cell = worksheet.cell(
                        row=row_index,
                        column=column_index,
                        value=value,
                    )
                    cell.alignment = center
                    cell.border = table_border
                    cell.number_format = (
                        "0.000" if is_numeric else "@"
                    )
                    cell.comment = Comment(
                        comment,
                        "batch_experiment",
                    )

            for column_index in range(2, last_column + 1):
                worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = 12

            data_range = (
                f"B5:{get_column_letter(last_column)}"
                f"{4 + len(q0_values)}"
            )
            worksheet.conditional_formatting.add(
                data_range,
                ColorScaleRule(
                    start_type="min",
                    start_color="FEE2E2",
                    mid_type="num",
                    mid_value=0,
                    mid_color="FEF3C7",
                    end_type="max",
                    end_color="DCFCE7",
                ),
            )

        if include_raw_data:
            raw_name = _excel_sheet_name(
                "원본결과",
                used_sheet_names,
            )
            raw_sheet = workbook.create_sheet(raw_name)
            raw_sheet.sheet_view.showGridLines = False
            raw_sheet.freeze_panes = "A2"
            raw_sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(SUMMARY_FIELDS))}"
                f"{len(model_rows) + 1}"
            )

            for column_index, field in enumerate(
                SUMMARY_FIELDS,
                start=1,
            ):
                cell = raw_sheet.cell(
                    row=1,
                    column=column_index,
                    value=field,
                )
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center
                cell.border = table_border

            text_fields = {
                "model",
                "model_path",
                "detail_json",
                "timestamp_utc",
                "run_name",
                "status",
                "selected_direction",
                "axis_source",
                "slice_bands_json",
                "error_type",
                "error_message",
                "traceback",
            }
            for row_index, row in enumerate(model_rows, start=2):
                for column_index, field in enumerate(
                    SUMMARY_FIELDS,
                    start=1,
                ):
                    raw_value: Any = row.get(field, "")
                    numeric_value = _float_or_none(raw_value)
                    if (
                        numeric_value is not None
                        and field not in text_fields
                    ):
                        raw_value = numeric_value
                    cell = raw_sheet.cell(
                        row=row_index,
                        column=column_index,
                        value=raw_value,
                    )
                    cell.border = table_border
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=False,
                    )

            for column_index, field in enumerate(
                SUMMARY_FIELDS,
                start=1,
            ):
                width = 14
                if field in {
                    "model",
                    "run_name",
                    "selected_direction",
                    "status",
                }:
                    width = 18
                elif field in {
                    "model_path",
                    "detail_json",
                    "error_message",
                    "slice_bands_json",
                }:
                    width = 40
                elif field == "traceback":
                    width = 48
                raw_sheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = width

        output_name = _excel_report_filename(
            model_path,
            used_filenames,
        )
        output_path = output_dir / output_name
        workbook.save(output_path)
        output_paths.append(output_path)
        print(
            f"모델별 Excel: {output_path} "
            f"(밴드 시트 {maximum_band_count}개)"
        )

    return output_paths


def run_batch(
    batch_config_path: str | Path,
    cli_overrides: Mapping[str, Any] | None = None,
) -> int:
    batch, resolved_batch_path = load_batch_config(batch_config_path)
    if cli_overrides:
        for key, value in cli_overrides.items():
            if value is not None:
                batch[key] = value

    models = discover_stl_files(
        Path(batch["models_dir"]),
        bool(batch["recursive"]),
    )
    configurations = expand_experiment_configs(
        Path(batch["base_config_file"]),
        batch["base_overrides"],
        batch["parameter_grid"],
        batch["parameter_sets"],
    )

    run_name = str(batch["run_name"] or _default_run_name())
    run_dir = Path(batch["output_dir"]) / _safe_filename(run_name)
    details_dir = run_dir / "details"
    run_dir.mkdir(parents=True, exist_ok=False)
    if batch["save_detailed_json"]:
        details_dir.mkdir(parents=True, exist_ok=True)

    csv_path = run_dir / "summary.csv"
    jsonl_path = run_dir / "summary.jsonl"
    manifest_path = run_dir / "manifest.json"
    total_planned = (
        len(models)
        * len(configurations)
        * int(batch["repeats"])
    )
    manifest: dict[str, Any] = {
        "status": "running",
        "started_utc": _utc_now(),
        "finished_utc": None,
        "batch_config_path": str(resolved_batch_path),
        "run_name": run_name,
        "models_dir": str(Path(batch["models_dir"]).resolve()),
        "models": [str(path.resolve()) for path in models],
        "configuration_count": len(configurations),
        "configurations": configurations,
        "repeats": int(batch["repeats"]),
        "planned_execution_count": total_planned,
        "completed_execution_count": 0,
        "success_count": 0,
        "error_count": 0,
        "summary_csv": str(csv_path.resolve()),
        "summary_jsonl": str(jsonl_path.resolve()),
        "model_excel_files": [],
        "excel_report_error": None,
    }
    _write_manifest(manifest_path, manifest)

    print(
        f"모델 {len(models)}개 x 설정 {len(configurations)}개 "
        f"x 반복 {batch['repeats']}회"
    )
    print(f"총 실행 예정: {total_planned}회")
    print(f"결과 폴더: {run_dir}")

    interrupted = False
    experiment_index = 0
    with csv_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as csv_stream, jsonl_path.open(
        "w",
        encoding="utf-8",
        newline="\n",
    ) as jsonl_stream:
        writer = csv.DictWriter(
            csv_stream,
            fieldnames=SUMMARY_FIELDS,
            extrasaction="ignore",
        )
        writer.writeheader()
        csv_stream.flush()

        try:
            for model_path in models:
                for configuration_index, config in enumerate(
                    configurations,
                    start=1,
                ):
                    for repeat in range(
                        1,
                        int(batch["repeats"]) + 1,
                    ):
                        experiment_index += 1
                        print(
                            f"[{experiment_index}/{total_planned}] "
                            f"{model_path.name} "
                            f"config={configuration_index}/"
                            f"{len(configurations)} "
                            f"repeat={repeat}"
                        )
                        started = time.perf_counter()
                        detail_path: Path | None = None

                        try:
                            result = analyze_stl(model_path, config)
                            wall_runtime = (
                                time.perf_counter() - started
                            )
                            if batch["save_detailed_json"]:
                                stem = _safe_filename(
                                    model_path.stem
                                )
                                detail_path = details_dir / (
                                    f"{experiment_index:06d}_{stem}_"
                                    f"c{configuration_index:04d}_"
                                    f"r{repeat:03d}.json"
                                )
                                write_json(result, detail_path)

                            row = result_to_summary(
                                result,
                                run_name=run_name,
                                experiment_index=experiment_index,
                                configuration_index=(
                                    configuration_index
                                ),
                                repeat=repeat,
                                model_path=model_path,
                                detail_json=detail_path,
                                wall_runtime_seconds=wall_runtime,
                                config=config,
                            )
                            manifest["success_count"] += 1
                            print(
                                f"  -> {row['selected_direction']} "
                                f"{float(row['selected_angle']):.6g} "
                                f"deg, "
                                f"bands={row['slice_band_count']}, "
                                f"loss={float(row['selected_loss']):.6g}"
                            )
                        except (
                            ConfigError,
                            MeshAnalysisError,
                            OSError,
                            ValueError,
                        ) as exc:
                            wall_runtime = (
                                time.perf_counter() - started
                            )
                            row = error_to_summary(
                                exc,
                                run_name=run_name,
                                experiment_index=experiment_index,
                                configuration_index=(
                                    configuration_index
                                ),
                                repeat=repeat,
                                model_path=model_path,
                                wall_runtime_seconds=wall_runtime,
                                config=config,
                            )
                            manifest["error_count"] += 1
                            print(
                                f"  -> 오류: {type(exc).__name__}: "
                                f"{exc}",
                                file=sys.stderr,
                            )
                            if not batch["continue_on_error"]:
                                writer.writerow(row)
                                csv_stream.flush()
                                _write_jsonl_row(
                                    jsonl_stream,
                                    row,
                                )
                                manifest[
                                    "completed_execution_count"
                                ] = experiment_index
                                raise

                        writer.writerow(row)
                        csv_stream.flush()
                        _write_jsonl_row(jsonl_stream, row)
                        manifest[
                            "completed_execution_count"
                        ] = experiment_index
                        _write_manifest(
                            manifest_path,
                            manifest,
                        )
        except KeyboardInterrupt:
            interrupted = True
            print(
                "\n사용자 요청으로 중단했습니다. 완료된 결과는 "
                "그대로 보존됩니다.",
                file=sys.stderr,
            )

    excel_report_failed = False
    if batch["save_model_excel"]:
        try:
            excel_paths = write_model_excel_reports(
                csv_path,
                run_dir / "model_excels",
                measurements=batch["excel_measurements"],
                include_raw_data=bool(
                    batch["excel_include_raw_data"]
                ),
            )
            manifest["model_excel_files"] = [
                str(path.resolve()) for path in excel_paths
            ]
        except (BatchConfigError, OSError, ValueError) as exc:
            excel_report_failed = True
            manifest["excel_report_error"] = (
                f"{type(exc).__name__}: {exc}"
            )
            print(
                f"Excel 보고서 생성 오류: {exc}",
                file=sys.stderr,
            )

    manifest["status"] = (
        "interrupted" if interrupted else "completed"
    )
    if excel_report_failed and not interrupted:
        manifest["status"] = "completed_with_excel_error"
    manifest["finished_utc"] = _utc_now()
    _write_manifest(manifest_path, manifest)

    print(
        f"완료: 성공 {manifest['success_count']}회, "
        f"오류 {manifest['error_count']}회"
    )
    print(f"CSV: {csv_path}")
    print(f"JSONL: {jsonl_path}")
    if manifest["model_excel_files"]:
        print(f"모델별 Excel 폴더: {run_dir / 'model_excels'}")

    if interrupted:
        return 130
    if manifest["error_count"] or excel_report_failed:
        return 1
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "폴더의 STL 파일을 설정 조합별로 반복 분석하고 "
            "모델별 Excel과 밴드별 시트를 생성합니다."
        )
    )
    parser.add_argument(
        "--config",
        default="batch_config.json",
        help="배치 설정 JSON 경로 (기본: batch_config.json)",
    )
    parser.add_argument(
        "--models-dir",
        help="배치 설정의 models_dir 덮어쓰기",
    )
    parser.add_argument(
        "--output-dir",
        help="배치 설정의 output_dir 덮어쓰기",
    )
    parser.add_argument(
        "--run-name",
        help="결과 하위 폴더 이름",
    )
    parser.add_argument(
        "--repeats",
        type=int,
        help="동일 설정 반복 횟수",
    )
    parser.add_argument(
        "--recursive",
        action=argparse.BooleanOptionalAction,
        default=None,
        help="하위 폴더의 STL까지 검색",
    )
    parser.add_argument(
        "--save-detailed-json",
        action=argparse.BooleanOptionalAction,
        default=None,
        help="실행별 전체 JSON 저장",
    )
    parser.add_argument(
        "--continue-on-error",
        action=argparse.BooleanOptionalAction,
        default=None,
        help="한 모델이 실패해도 다음 실행 계속",
    )
    parser.add_argument(
        "--save-model-excel",
        action=argparse.BooleanOptionalAction,
        default=None,
        help="모델별 Excel 파일 생성",
    )
    parser.add_argument(
        "--excel-include-raw-data",
        action=argparse.BooleanOptionalAction,
        default=None,
        help="모델별 Excel에 원본결과 시트 포함",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(
        list(argv) if argv is not None else None
    )
    overrides: dict[str, Any] = {
        "models_dir": (
            None
            if args.models_dir is None
            else Path(args.models_dir).resolve()
        ),
        "output_dir": (
            None
            if args.output_dir is None
            else Path(args.output_dir).resolve()
        ),
        "run_name": args.run_name,
        "repeats": args.repeats,
        "recursive": args.recursive,
        "save_detailed_json": args.save_detailed_json,
        "continue_on_error": args.continue_on_error,
        "save_model_excel": args.save_model_excel,
        "excel_include_raw_data": args.excel_include_raw_data,
    }
    try:
        if args.repeats is not None and args.repeats <= 0:
            raise BatchConfigError(
                "--repeats는 1 이상이어야 합니다."
            )
        return run_batch(args.config, overrides)
    except (
        BatchConfigError,
        ConfigError,
        OSError,
        ValueError,
    ) as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())